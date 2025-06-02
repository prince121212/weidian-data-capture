#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
删除指定分类ID的行
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment

def delete_specified_categories():
    """删除指定分类ID的行"""
    
    # 要删除的分类ID列表
    delete_ids = [
        "124372511", "113657275", "135496663", "135496594", 
        "113791987", "113703833", "125170448", "115710964", "115026831"
    ]
    
    print("🗑️ 删除指定分类ID的行")
    print("=" * 50)
    
    try:
        # 读取原文件
        df = pd.read_excel('data/shop_categories_with_links.xlsx', sheet_name='分类汇总')
        print(f"📊 原文件总行数: {len(df)}")
        
        # 显示要删除的分类信息
        print(f"\n🎯 要删除的分类ID: {delete_ids}")
        print(f"\n📋 要删除的分类详情:")
        
        delete_rows = df[df['分类ID'].astype(str).isin(delete_ids)]
        
        if len(delete_rows) > 0:
            for i, row in delete_rows.iterrows():
                print(f"  - {row['分类名称']} (ID: {row['分类ID']}, 路径: {row['完整分类路径']})")
        else:
            print("  未找到要删除的分类")
        
        # 删除指定行
        df_filtered = df[~df['分类ID'].astype(str).isin(delete_ids)]
        
        print(f"\n📊 删除后行数: {len(df_filtered)}")
        print(f"📊 删除了 {len(df) - len(df_filtered)} 行")
        
        # 保存到新文件
        output_filename = 'data/shop_categories_with_links_filtered.xlsx'
        
        # 创建工作簿
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        
        # 创建分类汇总表
        summary_sheet = workbook.create_sheet('分类汇总')
        
        # 获取列名
        headers = df_filtered.columns.tolist()
        
        # 写入表头并设置样式
        for col, header in enumerate(headers, 1):
            cell = summary_sheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        # 写入数据
        for row_idx, (_, row) in enumerate(df_filtered.iterrows(), 2):
            for col_idx, header in enumerate(headers, 1):
                value = row[header]
                summary_sheet.cell(row=row_idx, column=col_idx, value=str(value))
        
        # 自动调整列宽
        for column in summary_sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)  # 最大宽度50
            summary_sheet.column_dimensions[column_letter].width = adjusted_width
        
        # 创建统计表
        stats_sheet = workbook.create_sheet('分类统计')
        
        # 统计信息
        stats_data = [
            ["统计项目", "数值"],
            ["原总分类数", len(df)],
            ["删除分类数", len(df) - len(df_filtered)],
            ["剩余分类数", len(df_filtered)],
            ["", ""],
            ["按层级统计", ""],
        ]
        
        # 按层级统计
        level_count = df_filtered['层级'].value_counts().sort_index()
        for level, count in level_count.items():
            stats_data.append([f"第{level}级分类", count])
        
        stats_data.extend([
            ["", ""],
            ["按类型统计", ""],
        ])
        
        # 按类型统计
        type_count = df_filtered['分类类型'].value_counts()
        for cat_type, count in type_count.items():
            stats_data.append([cat_type, count])
        
        stats_data.extend([
            ["", ""],
            ["删除的分类ID", ""],
        ])
        
        # 添加删除的分类ID
        for delete_id in delete_ids:
            stats_data.append([delete_id, "已删除"])
        
        # 写入统计数据
        for row, (key, value) in enumerate(stats_data, 1):
            stats_sheet.cell(row=row, column=1, value=key)
            stats_sheet.cell(row=row, column=2, value=value)
            
            # 设置表头样式
            if row == 1 or key in ["按层级统计", "按类型统计", "删除的分类ID"]:
                stats_sheet.cell(row=row, column=1).font = Font(bold=True)
                stats_sheet.cell(row=row, column=2).font = Font(bold=True)
        
        # 调整统计表列宽
        stats_sheet.column_dimensions['A'].width = 20
        stats_sheet.column_dimensions['B'].width = 15
        
        # 保存文件
        workbook.save(output_filename)
        print(f"\n✅ 过滤后的文件已保存到: {output_filename}")
        
        # 显示剩余分类的前10个
        print(f"\n📋 剩余分类示例 (前10个):")
        for i, row in df_filtered.head(10).iterrows():
            print(f"  {i+1}. {row['分类名称']} (ID: {row['分类ID']})")
        
        if len(df_filtered) > 10:
            print(f"  ... 还有 {len(df_filtered) - 10} 个分类")
        
        return True
        
    except Exception as e:
        print(f"❌ 处理失败: {e}")
        import traceback
        traceback.print_exc()
        return False

def backup_original_file():
    """备份原文件"""
    try:
        import shutil
        backup_filename = 'data/shop_categories_with_links_backup.xlsx'
        shutil.copy2('data/shop_categories_with_links.xlsx', backup_filename)
        print(f"📁 原文件已备份到: {backup_filename}")
        return True
    except Exception as e:
        print(f"❌ 备份失败: {e}")
        return False

def main():
    """主函数"""
    print("🗑️ 分类删除工具")
    print("=" * 50)
    
    # 备份原文件
    if backup_original_file():
        # 删除指定分类
        if delete_specified_categories():
            print(f"\n🎉 任务完成!")
            print(f"📁 原文件备份: data/shop_categories_with_links_backup.xlsx")
            print(f"📁 过滤后文件: data/shop_categories_with_links_filtered.xlsx")
        else:
            print(f"\n❌ 删除失败")
    else:
        print(f"\n❌ 备份失败，操作中止")

if __name__ == "__main__":
    main()
