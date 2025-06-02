#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
获取至尊青花分类的真实商品
通过搜索所有商品并筛选包含"至尊青花"的商品
"""

import requests
import json
import openpyxl
from openpyxl.styles import Font, Alignment
import time

class ZhizunQinghuaCollector:
    def __init__(self):
        """初始化商品搜集器"""
        self.products = []
        
    def get_all_products_and_filter(self, shop_id, target_keywords=["至尊青花", "至尊"]):
        """获取所有商品并筛选包含目标关键词的商品"""
        print(f"🔍 开始获取所有商品并筛选包含 {target_keywords} 的商品...")
        
        url = "https://thor.weidian.com/decorate/shopDetail.tab.getItemList/1.0"
        
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            "Referer": f"https://h5.weidian.com/decoration/shop-category/?userid={shop_id}",
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8"
        }
        
        all_products = []
        filtered_products = []
        page = 0
        max_pages = 20  # 最大页数限制
        
        while page < max_pages:
            offset = page * 20
            
            # 获取所有商品（不指定分类）
            api_params = {
                "shopId": shop_id,
                "sortOrder": "desc",
                "offset": offset,
                "limit": 20,
                "from": "h5",
                "showItemTag": True
            }
            
            params = {
                "param": json.dumps(api_params)
            }
            
            try:
                print(f"  📄 获取第 {page + 1} 页数据...")
                response = requests.get(url, params=params, headers=headers, timeout=15)
                
                if response.status_code == 200:
                    data = response.json()
                    
                    if data.get("status", {}).get("code") == 0:
                        result = data.get("result", {})
                        item_list = result.get("itemList", [])
                        
                        if not item_list:
                            print(f"  ✅ 第 {page + 1} 页无数据，获取完成")
                            break
                        
                        # 解析商品信息并筛选
                        for item in item_list:
                            product_info = self.parse_product_info(item)
                            if product_info:
                                all_products.append(product_info)
                                
                                # 检查是否包含目标关键词
                                item_name = product_info.get("商品名称", "")
                                for keyword in target_keywords:
                                    if keyword in item_name:
                                        filtered_products.append(product_info)
                                        print(f"  🎯 找到目标商品: {item_name}")
                                        break
                        
                        print(f"  ✅ 第 {page + 1} 页获取到 {len(item_list)} 个商品，筛选出 {len([p for p in filtered_products if any(k in p.get('商品名称', '') for k in target_keywords)])} 个目标商品")
                        
                        # 如果返回的商品数量少于20，说明已经是最后一页
                        if len(item_list) < 20:
                            print(f"  ✅ 已获取完所有商品")
                            break
                            
                    else:
                        print(f"  ❌ API返回错误: {data.get('status', {}).get('message', '未知错误')}")
                        break
                else:
                    print(f"  ❌ API请求失败，状态码: {response.status_code}")
                    break
                    
            except Exception as e:
                print(f"  ❌ 请求异常: {e}")
                break
            
            page += 1
            time.sleep(1)  # 避免请求过快
        
        print(f"\n📊 搜索完成:")
        print(f"  总商品数: {len(all_products)}")
        print(f"  目标商品数: {len(filtered_products)}")
        
        return filtered_products
    
    def parse_product_info(self, item):
        """解析单个商品信息"""
        try:
            product_info = {
                "商品ID": item.get("itemId", ""),
                "商品名称": item.get("itemName", ""),
                "价格": item.get("price", ""),
                "销量": item.get("soldNum", 0),
                "库存": item.get("stock", 0),
                "商品链接": f"https://weidian.com/item.html?itemID={item.get('itemId', '')}" if item.get('itemId') else "",
                "图片链接": item.get("thumbs", [""])[0] if item.get("thumbs") else "",
                "商品描述": item.get("itemDesc", ""),
                "店铺ID": item.get("shopId", ""),
                "创建时间": item.get("addTime", ""),
                "更新时间": item.get("updateTime", "")
            }
            
            return product_info
            
        except Exception as e:
            print(f"  ⚠️ 解析商品信息失败: {e}")
            return None
    
    def save_products_to_xlsx(self, products, category_name, filename):
        """保存商品信息到xlsx文件"""
        try:
            print(f"\n💾 正在保存商品数据到 {filename}...")
            
            # 创建工作簿
            workbook = openpyxl.Workbook()
            
            # 删除默认工作表
            workbook.remove(workbook.active)
            
            # 创建商品列表表
            products_sheet = workbook.create_sheet(f'{category_name}_商品列表')
            
            # 设置表头
            headers = ["商品ID", "商品名称", "价格", "销量", "库存", "商品链接", "图片链接", "商品描述", "店铺ID", "创建时间", "更新时间"]
            
            # 写入表头并设置样式
            for col, header in enumerate(headers, 1):
                cell = products_sheet.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
            
            # 写入商品数据
            for row, product in enumerate(products, 2):
                for col, header in enumerate(headers, 1):
                    value = product.get(header, "")
                    products_sheet.cell(row=row, column=col, value=str(value))
            
            # 自动调整列宽
            for column in products_sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)  # 最大宽度50
                products_sheet.column_dimensions[column_letter].width = adjusted_width
            
            # 创建统计表
            stats_sheet = workbook.create_sheet('商品统计')
            
            # 统计信息
            stats_data = [
                ["统计项目", "数值"],
                ["分类名称", category_name],
                ["总商品数", len(products)],
                ["", ""],
                ["价格统计", ""],
            ]
            
            # 价格统计
            if products:
                prices = []
                for product in products:
                    try:
                        price_str = str(product.get("价格", "0")).replace("¥", "").replace(",", "")
                        price = float(price_str) if price_str else 0
                        if price > 0:
                            prices.append(price)
                    except:
                        continue
                
                if prices:
                    stats_data.extend([
                        ["最高价格", f"¥{max(prices):.2f}"],
                        ["最低价格", f"¥{min(prices):.2f}"],
                        ["平均价格", f"¥{sum(prices)/len(prices):.2f}"],
                    ])
            
            stats_data.extend([
                ["", ""],
                ["销量统计", ""],
            ])
            
            # 销量统计
            if products:
                sales = [int(str(product.get("销量", "0"))) for product in products]
                sales = [s for s in sales if s >= 0]
                
                if sales:
                    stats_data.extend([
                        ["总销量", sum(sales)],
                        ["最高销量", max(sales)],
                        ["平均销量", f"{sum(sales)/len(sales):.1f}"],
                    ])
            
            # 写入统计数据
            for row, (key, value) in enumerate(stats_data, 1):
                stats_sheet.cell(row=row, column=1, value=key)
                stats_sheet.cell(row=row, column=2, value=value)
                
                # 设置表头样式
                if row == 1 or key in ["价格统计", "销量统计"]:
                    stats_sheet.cell(row=row, column=1).font = Font(bold=True)
                    stats_sheet.cell(row=row, column=2).font = Font(bold=True)
            
            # 调整统计表列宽
            stats_sheet.column_dimensions['A'].width = 20
            stats_sheet.column_dimensions['B'].width = 20
            
            # 保存文件
            workbook.save(filename)
            print(f"✅ 商品数据已成功保存到: {filename}")
            
            return True
            
        except Exception as e:
            print(f"❌ 保存xlsx文件失败: {e}")
            import traceback
            traceback.print_exc()
            return False


def main():
    """主函数"""
    print("=" * 60)
    print("🎯 至尊青花商品搜集工具")
    print("=" * 60)
    
    collector = ZhizunQinghuaCollector()
    
    # 获取商品信息
    shop_id = "1286456178"
    target_keywords = ["至尊青花", "至尊"]
    
    products = collector.get_all_products_and_filter(shop_id, target_keywords)
    
    if products:
        # 保存到xlsx文件
        filename = "data/至尊青花_商品列表.xlsx"
        success = collector.save_products_to_xlsx(products, "至尊青花", filename)
        
        if success:
            print(f"\n🎉 任务完成!")
            print(f"📁 文件已保存: {filename}")
            print(f"📊 包含内容:")
            print(f"  - 至尊青花_商品列表: {len(products)}个商品的详细信息")
            print(f"  - 商品统计: 价格和销量统计信息")
            
            # 显示所有商品
            print(f"\n🎯 找到的至尊青花商品:")
            for i, product in enumerate(products):
                print(f"  {i+1}. {product['商品名称']} (ID: {product['商品ID']}, 价格: {product['价格']})")
        else:
            print(f"\n❌ 保存文件失败")
    else:
        print(f"\n❌ 未能获取到任何至尊青花相关商品")


if __name__ == "__main__":
    main()
