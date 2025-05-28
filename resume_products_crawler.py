import xlrd
import xlwt
import time
import os
from openpyxl import load_workbook, Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# 假设 single_product_crawler.py 中的函数与 batch_products_crawler.py 中使用的相同
from single_product_crawler import extract_item_id, fetch_detail_page, parse_detail_page_html, extract_detail_api, extract_all_images, extract_with_selenium

def resume_batch_crawl(input_xls, url_column, output_xls, start_row_one_indexed, delay=2.0, use_selenium=True):
    """
    从指定行开始，批量爬取商品详情，并追加到输出文件。

    参数:
    - input_xls: 输入的商品列表Excel文件
    - url_column: 包含商品链接的列名
    - output_xls: 输出的Excel文件
    - start_row_one_indexed: 从第几行开始爬取 (1-indexed, 指的是数据行，不包括表头)
    - delay: 每个请求之间的延迟时间(秒)
    - use_selenium: 是否使用Selenium提取更多内容(速度较慢但信息更完整)
    """
    # 创建输出目录
    os.makedirs(os.path.dirname(output_xls), exist_ok=True)

    print(f"正在读取输入文件: {input_xls}")
    wb_input = xlrd.open_workbook(input_xls)
    sheet_input = wb_input.sheet_by_index(0)
    headers_input = [sheet_input.cell_value(0, col) for col in range(sheet_input.ncols)]

    if url_column not in headers_input:
        print(f"错误: 输入文件中未找到列 '{url_column}'")
        return

    url_idx = headers_input.index(url_column)
    
    # 确保 start_row_one_indexed 是有效的
    if start_row_one_indexed <= 0:
        print(f"错误: start_row_one_indexed ({start_row_one_indexed}) 必须大于 0。")
        return
    if start_row_one_indexed > sheet_input.nrows -1: # sheet_input.nrows 包括表头
        print(f"指定的开始行 ({start_row_one_indexed}) 超出文件总行数 ({sheet_input.nrows - 1})。无需爬取。")
        return

    # 调整为 0-indexed 的数据行号，并加上表头行
    actual_start_row_xlrd = start_row_one_indexed 


    print(f"将从输入文件的第 {start_row_one_indexed} 条商品数据开始爬取 (对应 Excel 表格中的第 {actual_start_row_xlrd + 1} 行)。")

    new_rows_data = []
    all_new_keys = set()

    # 临时文件，用于保存中间结果
    temp_output_suffix = f"_temp_from_row_{start_row_one_indexed}.xlsx"
    temp_output = output_xls.replace(".xlsx", temp_output_suffix).replace(".xls", temp_output_suffix)
    if not temp_output.endswith(".xlsx"): # 确保是 .xlsx
        temp_output += ".xlsx"


    for row_idx_xlrd in range(actual_start_row_xlrd, sheet_input.nrows):
        current_item_number = row_idx_xlrd # 因为 actual_start_row_xlrd 已经是基于表头的行号
        total_items_to_scan = sheet_input.nrows -1 # 总数据行数

        url = sheet_input.cell_value(row_idx_xlrd, url_idx)
        if not url or not isinstance(url, str) or not url.startswith('http'):
            print(f"跳过无效链接: 输入文件第 {row_idx_xlrd + 1} 行")
            continue

        print(f"[{current_item_number - actual_start_row_xlrd + 1}/{total_items_to_scan - actual_start_row_xlrd + 1}] (总进度: {current_item_number}/{total_items_to_scan}) 正在爬取: {url}")

        try:
            item_id = extract_item_id(url)
            html = fetch_detail_page(url)
            if not html:
                print(f"获取页面失败: {url}")
                continue

            html_details = parse_detail_page_html(html)
            images_from_html = extract_all_images(html)
            api_details = extract_detail_api(url)
            selenium_details = {}
            selenium_images = []
            if use_selenium:
                print("使用Selenium提取更多内容...")
                selenium_details, selenium_images = extract_with_selenium(url)

            all_details = {**html_details, **api_details, **selenium_details}
            all_images = list(set(images_from_html + selenium_images))

            print(f"  - 提取到 {len(all_details)} 个详情字段")
            print(f"  - 提取到 {len(all_images)} 张图片链接")

            row_dict = {'商品ID': item_id}
            row_dict.update(all_details)
            row_dict['图片链接'] = '\n'.join(all_images)
            
            new_rows_data.append(row_dict)
            all_new_keys.update(all_details.keys())

            if len(new_rows_data) % 5 == 0:
                save_to_excel_openpyxl(new_rows_data, all_new_keys, temp_output, append=True, input_headers_order=headers_input, url_column=url_column)
                print(f"已保存临时结果到: {temp_output}")

            print(f"等待 {delay} 秒...")
            time.sleep(delay)

        except Exception as e:
            print(f"处理商品失败: {url}, 错误: {e}")
            import traceback
            traceback.print_exc()

    if new_rows_data:
        print(f"所有新商品爬取完成，准备合并到最终文件: {output_xls}")
        save_to_excel_openpyxl(new_rows_data, all_new_keys, output_xls, append=True, input_headers_order=headers_input, url_column=url_column)
        print(f"已追加新的商品详情到: {output_xls}")
        if os.path.exists(temp_output):
            try:
                os.remove(temp_output)
                print(f"已删除临时文件: {temp_output}")
            except Exception as e:
                print(f"删除临时文件失败: {e}")
    else:
        print("没有新的商品数据被爬取。")


def save_to_excel_openpyxl(rows_data, data_keys, output_file, append=False, input_headers_order=None, url_column=None):
    """使用 openpyxl 保存或追加数据到 Excel (.xlsx) 文件"""
    
    # 定义期望的表头顺序，基于原脚本的 save_to_xls 和输入文件的表头
    # '商品ID' 和 '图片链接' 是固定的
    base_fixed_keys = ['商品ID']
    # 尝试从 data_keys 中获取原脚本中定义的其他固定字段
    # 注意：这些键可能不存在于当前批次爬取到的 all_new_keys 中
    original_script_fixed_keys = ['商品标题', '价格', '销量', '详细描述'] 
    
    # 合并所有可能的表头字段，并去重，同时尽量保持顺序
    # 1. 商品ID
    # 2. input_headers_order 中除了 '商品ID' 和 '图片链接' (如果存在) 之外的原始列
    # 3. data_keys 中新增的，不在 input_headers_order 中的字段
    # 4. 图片链接
    
    current_dynamic_keys = sorted([k for k in data_keys if k not in base_fixed_keys])

    # 构建最终表头
    # 如果 output_file 已存在，我们会尝试读取它的表头以保持一致性
    # 否则，我们会基于 input_headers_order 和新抓取的 keys 来创建表头
    
    final_headers = []
    existing_workbook = None
    sheet = None
    start_write_row = 1 # 1-indexed, for header

    if append and os.path.exists(output_file):
        try:
            existing_workbook = load_workbook(output_file)
            sheet = existing_workbook.active
            final_headers = [cell.value for cell in sheet[1]] # 读取第一行作为表头
            start_write_row = sheet.max_row + 1
            print(f"正在追加到现有文件 {output_file}。从第 {start_write_row} 行开始写入。")
            print(f"检测到现有表头: {final_headers}")

            # 检查新数据中是否有 기존表头中没有的字段，并添加它们 (这可能导致列不一致，需要用户注意)
            new_unique_keys_to_add_to_header = [k for k in (base_fixed_keys + current_dynamic_keys + ['图片链接']) if k not in final_headers]
            if new_unique_keys_to_add_to_header:
                print(f"警告: 新爬取的数据包含现有输出文件中没有的列: {new_unique_keys_to_add_to_header}。这些列将被添加到表头的末尾。")
                for new_key in new_unique_keys_to_add_to_header:
                    final_headers.append(new_key)
                    sheet.cell(row=1, column=len(final_headers), value=new_key) # 更新表头行
        
        except Exception as e:
            print(f"读取现有输出文件 {output_file} 失败: {e}。将创建一个新文件或覆盖。")
            existing_workbook = Workbook() # 创建新的工作簿
            sheet = existing_workbook.active
            # 清空 final_headers 以便重新生成
            final_headers = [] 
            start_write_row = 1
    else:
        existing_workbook = Workbook()
        sheet = existing_workbook.active
        start_write_row = 1
        print(f"创建新文件或覆盖: {output_file}")

    if not final_headers: # 如果文件是新的，或者读取表头失败，则生成表头
        # 优先使用 input_xls 的表头顺序作为基础
        if input_headers_order:
            temp_headers = [h for h in input_headers_order if (url_column is None or h != url_column) and h not in ['商品ID', '图片链接']] # 排除爬虫自身会处理的列
        else:
            temp_headers = []

        # 确保 '商品ID' 在最前面
        final_headers.append('商品ID')
        
        # 添加原脚本中的固定字段 (如果它们在当前抓取到的数据中)
        for key in original_script_fixed_keys:
            if key in data_keys and key not in final_headers:
                final_headers.append(key)
        
        # 添加 input_xls 中的其他表头字段
        for key in temp_headers:
            if key not in final_headers: # 避免重复
                 final_headers.append(key)

        # 添加本次爬取到的其他动态字段
        other_dynamic_keys = sorted([k for k in data_keys if k not in final_headers and k not in base_fixed_keys])
        final_headers.extend(other_dynamic_keys)
        
        # 确保 '图片链接' 在最后面
        if '图片链接' in final_headers: # 如果已存在，先移除
            final_headers.remove('图片链接')
        final_headers.append('图片链接')
        
        print(f"生成新的表头: {final_headers}")
        if start_write_row == 1: # 只有在文件是新创建或表头行为空时才写入表头
             for col_idx, header_title in enumerate(final_headers, 1):
                sheet.cell(row=1, column=col_idx, value=header_title)
        if start_write_row == 1 and rows_data: # 如果写入了表头，数据从第二行开始
            start_write_row = 2


    # 写入数据
    for row_dict in rows_data:
        for col_idx, header_key in enumerate(final_headers, 1):
            value = row_dict.get(header_key, '')
            # openpyxl 会自动处理多数类型，但确保字符串化复杂对象
            if not isinstance(value, (str, int, float, bool)) and value is not None:
                value = str(value)
            elif value is None:
                value = ''
            sheet.cell(row=start_write_row, column=col_idx, value=value)
        start_write_row += 1

    try:
        existing_workbook.save(output_file)
    except Exception as e:
        print(f"保存Excel文件失败: {output_file}, 错误: {e}")
        # 尝试备用保存 (例如，如果文件被占用)
        try:
            backup_name = output_file.replace(".xlsx", f"_backup_{int(time.time())}.xlsx")
            existing_workbook.save(backup_name)
            print(f"已将结果保存到备用文件: {backup_name}")
        except Exception as e2:
            print(f"保存到备用文件也失败: {e2}")


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='从指定行开始批量爬取商品详情并追加到输出文件')
    parser.add_argument('--input', default='data/items_all.xls', help='输入商品列表xls/xlsx')
    parser.add_argument('--column', default='商品详情页链接', help='商品详情页链接列名')
    parser.add_argument('--output', default='data/items_all_details.xlsx', help='输出xlsx文件名 (推荐 .xlsx 以支持追加)')
    parser.add_argument('--start_row', type=int, required=True, help='从第几条商品数据开始爬取 (1-indexed, 不包括表头)')
    parser.add_argument('--delay', type=float, default=2.0, help='每个商品间隔秒数')
    parser.add_argument('--no-selenium', action='store_true', help='不使用Selenium（更快但信息可能不完整）')
    
    args = parser.parse_args()

    if args.start_row < 1:
        print("错误: --start_row 必须是大于等于1的整数。")
    else:
        # output 文件推荐使用 .xlsx 以更好地支持追加和处理大文件
        output_file = args.output
        if not output_file.endswith('.xlsx'):
            print(f"警告: 输出文件 '{output_file}' 不是 .xlsx 格式。建议使用 .xlsx 以获得更好的追加和兼容性。正在尝试继续...")

        resume_batch_crawl(
            args.input, 
            args.column, 
            output_file, 
            args.start_row, # start_row_one_indexed for data row
            args.delay, 
            use_selenium=not args.no_selenium
        ) 