#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
获取釉中青花玲珑分类的所有商品
"""

import requests
import json
import openpyxl
from openpyxl.styles import Font, Alignment
import time

class LinlongProductsCollector:
    def __init__(self):
        """初始化商品搜集器"""
        self.products = []
        
    def get_category_products_correct_api(self, shop_id, cate_id, category_name):
        """使用正确的API获取分类商品"""
        print(f"🔍 使用正确API获取分类 '{category_name}' 的商品...")
        print(f"  店铺ID: {shop_id}")
        print(f"  分类ID: {cate_id}")
        
        # 正确的API端点
        url = "https://thor.weidian.com/decorate/itemCate.getCateItemList/1.0"
        
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36 Edg/136.0.0.0",
            "Accept": "application/json, */*",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Accept-Encoding": "gzip, deflate, br, zstd",
            "Referer": "https://h5.weidian.com/",
            "Origin": "https://h5.weidian.com",
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-site"
        }
        
        all_products = []
        page = 0
        max_pages = 10  # 最大页数限制
        
        while page < max_pages:
            offset = page * 20
            
            # 构建正确的API参数
            api_params = {
                "cateId": str(cate_id),
                "shopId": str(shop_id),
                "offset": offset,
                "limit": 20,
                "sortField": "all",
                "sortType": "desc",
                "isQdFx": False,
                "isHideSold": False,
                "hideItemRealAmount": False,
                "from": "h5",
                "fanSpreadMode": -1,
                "isShopItemListConfOpen": False,
                "attrQuery": [],
                "isStockDown": 0,
                "isConsumerProtect": False,
                "hideItemComment": False
            }
            
            params = {
                "param": json.dumps(api_params, separators=(',', ':'))
            }
            
            try:
                print(f"  📄 获取第 {page + 1} 页数据...")
                response = requests.get(url, params=params, headers=headers, timeout=15)
                
                if response.status_code == 200:
                    data = response.json()
                    
                    if data.get("status", {}).get("code") == 0:
                        result = data.get("result", {})
                        item_list = result.get("itemList", [])
                        has_data = result.get("hasData", False)
                        
                        if not item_list or not has_data:
                            print(f"  ✅ 第 {page + 1} 页无数据，获取完成")
                            break
                        
                        # 解析商品信息
                        for item in item_list:
                            product_info = self.parse_product_info(item)
                            if product_info:
                                all_products.append(product_info)
                                print(f"    📦 {product_info['商品名称']} (ID: {product_info['商品ID']}, 价格: ¥{product_info['价格']})")
                        
                        print(f"  ✅ 第 {page + 1} 页获取到 {len(item_list)} 个商品")
                        
                        # 如果返回的商品数量少于20，说明已经是最后一页
                        if len(item_list) < 20:
                            print(f"  ✅ 已获取完所有商品")
                            break
                            
                    else:
                        error_msg = data.get('status', {}).get('message', '未知错误')
                        print(f"  ❌ API返回错误: {error_msg}")
                        break
                else:
                    print(f"  ❌ API请求失败，状态码: {response.status_code}")
                    print(f"  响应内容: {response.text[:200]}...")
                    break
                    
            except Exception as e:
                print(f"  ❌ 请求异常: {e}")
                break
            
            page += 1
            time.sleep(1)  # 避免请求过快
        
        print(f"\n📊 获取完成，共找到 {len(all_products)} 个商品")
        return all_products
    
    def parse_product_info(self, item):
        """解析单个商品信息"""
        try:
            product_info = {
                "商品ID": item.get("itemId", ""),
                "商品名称": item.get("itemName", ""),
                "价格": item.get("price", ""),
                "销量": item.get("sold", ""),
                "库存": item.get("stock", 0),
                "商品链接": item.get("itemUrl", ""),
                "图片链接": item.get("itemImg", ""),
                "商品描述": item.get("itemComment", ""),
                "原价": item.get("originalPrice", ""),
                "状态": item.get("status", ""),
                "是否有SKU": item.get("hasSku", False),
                "是否预售": item.get("preSale", False)
            }
            
            return product_info
            
        except Exception as e:
            print(f"  ⚠️ 解析商品信息失败: {e}")
            return None
    
    def save_products_to_xlsx(self, products, category_name, filename):
        """保存商品信息到xlsx文件"""
        try:
            print(f"\n💾 正在保存商品数据到 {filename}...")
            
            # 创建工作簿
            workbook = openpyxl.Workbook()
            
            # 删除默认工作表
            workbook.remove(workbook.active)
            
            # 创建商品列表表
            products_sheet = workbook.create_sheet(f'{category_name}_商品列表')
            
            # 设置表头
            headers = ["商品ID", "商品名称", "价格", "销量", "库存", "商品链接", "图片链接", "商品描述", "原价", "状态", "是否有SKU", "是否预售"]
            
            # 写入表头并设置样式
            for col, header in enumerate(headers, 1):
                cell = products_sheet.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
            
            # 写入商品数据
            for row, product in enumerate(products, 2):
                for col, header in enumerate(headers, 1):
                    value = product.get(header, "")
                    products_sheet.cell(row=row, column=col, value=str(value))
            
            # 自动调整列宽
            for column in products_sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)  # 最大宽度50
                products_sheet.column_dimensions[column_letter].width = adjusted_width
            
            # 创建统计表
            stats_sheet = workbook.create_sheet('商品统计')
            
            # 统计信息
            stats_data = [
                ["统计项目", "数值"],
                ["分类名称", category_name],
                ["总商品数", len(products)],
                ["", ""],
                ["商品详情", ""],
            ]
            
            # 添加每个商品的详细信息
            for i, product in enumerate(products, 1):
                stats_data.extend([
                    [f"商品{i}", ""],
                    ["  商品名称", product.get("商品名称", "")],
                    ["  商品ID", product.get("商品ID", "")],
                    ["  价格", f"¥{product.get('价格', '')}"],
                    ["  库存", product.get("库存", "")],
                    ["  商品链接", product.get("商品链接", "")],
                    ["", ""],
                ])
            
            # 写入统计数据
            for row, (key, value) in enumerate(stats_data, 1):
                stats_sheet.cell(row=row, column=1, value=key)
                stats_sheet.cell(row=row, column=2, value=value)
                
                # 设置表头样式
                if row == 1 or key in ["商品详情"] or key.startswith("商品") and not key.startswith("  "):
                    stats_sheet.cell(row=row, column=1).font = Font(bold=True)
                    stats_sheet.cell(row=row, column=2).font = Font(bold=True)
            
            # 调整统计表列宽
            stats_sheet.column_dimensions['A'].width = 20
            stats_sheet.column_dimensions['B'].width = 50
            
            # 保存文件
            workbook.save(filename)
            print(f"✅ 商品数据已成功保存到: {filename}")
            
            return True
            
        except Exception as e:
            print(f"❌ 保存xlsx文件失败: {e}")
            import traceback
            traceback.print_exc()
            return False


def main():
    """主函数"""
    print("=" * 60)
    print("🎯 釉中青花玲珑商品搜集工具")
    print("=" * 60)
    
    collector = LinlongProductsCollector()
    
    # 釉中青花玲珑分类信息
    shop_id = "1286456178"
    cate_id = "124372546"  # 釉中青花玲珑的分类ID
    category_name = "釉中青花玲珑"
    
    print(f"📋 目标分类信息:")
    print(f"  分类名称: {category_name}")
    print(f"  分类ID: {cate_id}")
    print(f"  店铺ID: {shop_id}")
    print(f"  预期商品数量: 6个")
    
    # 获取商品信息
    products = collector.get_category_products_correct_api(shop_id, cate_id, category_name)
    
    if products:
        # 保存到xlsx文件
        filename = "data/釉中青花玲珑_商品列表.xlsx"
        success = collector.save_products_to_xlsx(products, category_name, filename)
        
        if success:
            print(f"\n🎉 任务完成!")
            print(f"📁 文件已保存: {filename}")
            print(f"📊 包含内容:")
            print(f"  - {category_name}_商品列表: {len(products)}个商品的详细信息")
            print(f"  - 商品统计: 每个商品的详细信息")
            
            # 显示所有商品
            print(f"\n🎯 釉中青花玲珑分类的所有商品:")
            for i, product in enumerate(products, 1):
                print(f"  {i}. {product['商品名称']}")
                print(f"     ID: {product['商品ID']}")
                print(f"     价格: ¥{product['价格']}")
                print(f"     库存: {product['库存']}")
                print(f"     链接: {product['商品链接']}")
                print()
        else:
            print(f"\n❌ 保存文件失败")
    else:
        print(f"\n❌ 未能获取到任何商品信息")
        print(f"💡 可能的原因:")
        print(f"  1. 需要登录认证 (wdtoken)")
        print(f"  2. API参数需要调整")
        print(f"  3. 网络连接问题")


if __name__ == "__main__":
    main()
